"""Create an editable DOCX specification from repository sources. Requires python-docx."""
from pathlib import Path
import csv, re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
SOURCE=ROOT/'docs'/'erp'
OUT=ROOT.parent/'deliverables'
OUT.mkdir(exist_ok=True)

def keep(p): p.paragraph_format.keep_with_next=True
def paragraph(doc,text,style=None):
    p=doc.add_paragraph(style=style)
    # Keep filenames/identifiers as text, with ordinary readable punctuation.
    p.add_run(text.replace('`',''))
    return p

def add_markdown(doc,path):
    first_paragraph=len(doc.paragraphs)
    table_lines=[]
    def flush():
        nonlocal table_lines
        if not table_lines:return
        rows=[]
        for line in table_lines:
            cells=[c.strip().replace('`','') for c in line.strip('|').split('|')]
            if all(re.fullmatch(r'[: -]+',c) for c in cells):continue
            rows.append(cells)
        table_lines=[]
        if not rows:return
        table=doc.add_table(rows=1,cols=len(rows[0]));table.style='Light Shading Accent 1'
        for i,c in enumerate(rows[0]):table.rows[0].cells[i].text=c
        repeat=OxmlElement('w:tblHeader');table.rows[0]._tr.get_or_add_trPr().append(repeat)
        for row in rows[1:]:
            cells=table.add_row().cells
            for i,c in enumerate(row):cells[i].text=c
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    p.paragraph_format.space_after=Pt(5)
                    for r in p.runs:r.font.size=Pt(9)
        doc.add_paragraph()
    for line in path.read_text(encoding='utf-8').splitlines():
        if line.startswith('|'):table_lines.append(line);continue
        flush()
        if not line.strip():continue
        if line.startswith('# '):continue
        if line.startswith('## '):
            h=doc.add_heading(line[3:],level=1)
            if path.name=='TECHNICAL_SPECIFICATION.md' and re.match(r'\d+\.',line[3:]):h.paragraph_format.page_break_before=True
        elif line.startswith('### '):doc.add_heading(line[4:],level=2)
        elif line.startswith('- '):paragraph(doc,line[2:],'List Bullet')
        elif line.startswith('```'):continue
        else:paragraph(doc,line)
    flush()
    if path.name=='LOCAL_RUNBOOK.md':
        for p in doc.paragraphs[first_paragraph:]:
            if p.style.name=='Normal':p.paragraph_format.space_after=Pt(4)

def main():
    doc=Document();s=doc.sections[0]
    s.page_height=Inches(11.69);s.page_width=Inches(8.27)
    s.top_margin=Inches(.7);s.bottom_margin=Inches(.65);s.left_margin=Inches(.8);s.right_margin=Inches(.7)
    normal=doc.styles['Normal'];normal.font.name='Calibri';normal.font.size=Pt(10.5)
    normal.paragraph_format.space_after=Pt(7);normal.paragraph_format.line_spacing=1.08
    for name in ['Heading 1','Heading 2','Heading 3','Title','Subtitle']:
        doc.styles[name].font.name='Calibri';doc.styles[name].font.color.rgb=RGBColor.from_string('137D74')
    doc.styles['Heading 1'].font.size=Pt(20);doc.styles['Heading 2'].font.size=Pt(13)
    p=doc.add_paragraph();p.paragraph_format.space_before=Pt(50)
    r=p.add_run('FLOWPOS ERP');r.bold=True;r.font.size=Pt(38);r.font.color.rgb=RGBColor.from_string('12333C')
    p=doc.add_paragraph('Техническая спецификация\nи план разработки',style='Subtitle');p.runs[0].font.size=Pt(25)
    paragraph(doc,'Редакция 2.1  |  27 октября 2026 года')
    paragraph(doc,'Для владельца продукта, технического руководителя, бухгалтера-методолога, backend/frontend разработчиков и QA.')
    doc.add_paragraph()
    paragraph(doc,'Собственная браузерная бухгалтерия с торговлей, зарплатой, импортом, отчётностью и управляемым автопилотом.')
    paragraph(doc,'Первый реализованный срез: локальный синтетический стенд ЭСФ. Production, реальные подписи и отправки отключены.')
    paragraph(doc,'Пакет самодостаточен: архитектура, модель данных, API, единый backlog из 82 задач, критерии приёмки, исходники и фактические проверки.')
    doc.add_page_break();doc.add_heading('Карта документа',level=1)
    for title in ['Архитектура и границы реализации','Данные, учётное ядро, торговля и POS','ЭСФ: получение, отправка, статусы и подпись','Четыре портала, зарплата, отчётность и автопилот','Frontend, API, безопасность и эксплуатация','Тестирование, выпуск и production gates','Приложение A. Проверенная база репозитория','Приложение B. Источники и lineage','Приложение C. Приёмка и локальный запуск','Приложение D. Полный backlog']:
        paragraph(doc,title,'List Bullet')
    paragraph(doc,'Машиночитаемые OpenAPI, DDL и CSV находятся в developer ZIP; все ссылки внутри пакета относительные. Проектные контракты PostgreSQL и будущего ERP явно отделены от реализованного локального API.')
    doc.add_page_break();add_markdown(doc,SOURCE/'TECHNICAL_SPECIFICATION.md')
    for title,path in [('Приложение A. Проверенная база репозитория','REPOSITORY_BASELINE.md'),('Приложение B. Источники и lineage','SOURCES.md'),('Приложение C. Приёмка','ACCEPTANCE.md'),('Локальный запуск','LOCAL_RUNBOOK.md')]:
        doc.add_page_break();doc.add_heading(title,level=1);add_markdown(doc,SOURCE/path)
    doc.add_page_break();doc.add_heading('Приложение D. Полный backlog',level=1)
    paragraph(doc,'82 задачи с приоритетом, ролью, зависимостями и проверяемым результатом. DONE_LOCAL обозначает только локальное выполнение; все рабочие государственные подключения и весь остальной ERP остаются отдельными этапами.')
    rows=list(csv.DictReader((SOURCE/'BACKLOG.csv').open(encoding='utf-8-sig')))
    epic=None
    for r in rows:
        if r['epic']!=epic:
            epic=r['epic'];doc.add_heading(epic,level=2)
        h=paragraph(doc,f"{r['id']} · {r['title']}");keep(h);h.runs[0].bold=True
        p=paragraph(doc,f"{r['priority']}  |  {r['owner_role']}  |  Размер {r['size']}  |  {r['status']}  |  Зависимости: {r['dependencies'] or 'нет'}")
        p.paragraph_format.space_after=Pt(3);keep(p);p.runs[0].font.size=Pt(9);p.runs[0].font.color.rgb=RGBColor.from_string('527078')
        paragraph(doc,'Приёмка: '+r['acceptance']+'.')
    for title,path in [('Приложение E. Второй инженерный этап','STAGE2.md'),('История изменений','CHANGELOG.md')]:
        doc.add_page_break();doc.add_heading(title,level=1);add_markdown(doc,SOURCE/path)
    footer=s.footer.paragraphs[0];footer.alignment=2
    run=footer.add_run('FLOWPOS ERP  ·  ');run.font.size=Pt(8);run.font.color.rgb=RGBColor.from_string('6A7C82')
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
    doc.core_properties.title='FLOWPOS ERP — Техническая спецификация 2.1'
    doc.core_properties.subject='ERP architecture, ESF local stand, backlog and acceptance'
    doc.core_properties.author='FlowPOS';doc.core_properties.keywords='ERP, ESF, specification, synthetic stand'
    path=OUT/'FLOWPOS_ERP_Specification_v2_2026-10-27.docx';doc.save(path)
    print(path)

if __name__=='__main__':main()
